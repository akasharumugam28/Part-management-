import io

import pytest
from openpyxl import load_workbook
from werkzeug.http import parse_options_header

from conftest import create_part, upload


def filename(r):
    return parse_options_header(r.headers["Content-Disposition"])[1]["filename"]


@pytest.fixture
def parts(auth_client):
    create_part(auth_client, 1, {"image": upload()}, trolley_qty="40")
    create_part(auth_client, 2, part_name="Nut & <Washer>")
    return auth_client


def test_export_requires_login(client):
    r = client.get("/parts/1/export?format=xlsx&name=x")
    assert r.status_code == 302 and "/login" in r.location


def test_excel_download_has_only_that_part(parts):
    r = parts.get("/parts/1/export?format=xlsx&name=Line 3 stock")
    assert r.status_code == 200
    assert r.mimetype.endswith("spreadsheetml.sheet")
    assert filename(r) == "Line 3 stock.xlsx"
    ws = load_workbook(io.BytesIO(r.data)).active
    assert ws["A1"].value == "Line 3 stock"
    assert [c.value for c in ws[4]][1:5] == ["Image", "Trolley Qty", "Part Name", "Part Number"]
    assert ws["C5"].value == 40 and ws["E5"].value == "PN-0001" and ws["E6"].value is None
    assert len(ws._images) == 1


def test_pdf_download(parts):
    r = parts.get("/parts/2/export?format=pdf&name=report")
    assert r.status_code == 200
    assert r.mimetype == "application/pdf"
    assert r.data.startswith(b"%PDF")
    assert filename(r) == "report.pdf"


def test_missing_part_is_404(auth_client):
    assert auth_client.get("/parts/999/export?format=pdf&name=x").status_code == 404


@pytest.mark.parametrize("name,expected", [
    ("", "parts-list"),
    ("../../etc/passwd", "etcpasswd"),
    ("my file.xlsx", "my file"),
    ('a"b;c', "abc"),
])
def test_filename_is_sanitised(parts, name, expected):
    r = parts.get("/parts/1/export", query_string={"format": "pdf", "name": name})
    assert filename(r) == f"{expected}.pdf"


def test_bad_format_rejected(parts):
    assert parts.get("/parts/1/export?format=doc&name=x").status_code == 400


def test_download_button_sits_between_edit_and_delete(parts):
    html = parts.get("/").get_data(as_text=True)
    row = html[html.index('href="/parts/1/edit"'):]
    assert row.index('data-form-action="/parts/1/export"') < row.index('action="/parts/1/delete"')
    assert 'id="download-dialog"' in html and 'name="format" value="pdf"' in html
    page_head = html.split('class="page-head"')[1].split("</section>")[0]
    assert 'data-form-action="/parts/export"' in page_head and "Download All" in page_head
