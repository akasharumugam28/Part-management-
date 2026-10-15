import io
import sqlite3

import pytest
from openpyxl import load_workbook

from app import create_app
from conftest import create_part, db_rows, make_config


def counts(app, part_id=1):
    return tuple(db_rows(app, "SELECT total_trolleys, empty_trolleys FROM parts WHERE id = ?",
                         (part_id,))[0])


@pytest.mark.parametrize("total,empty,expected", [
    ("10", "4", (10, 4)), ("5", "5", (5, 5)), ("0", "0", (0, 0)), (" 8 ", "", (8, None)), ("", "", (None, None)),
])
def test_create_saves_counts(app, auth_client, total, empty, expected):
    r = create_part(auth_client, total_trolleys=total, empty_trolleys=empty)
    assert r.status_code == 303
    assert counts(app) == expected


@pytest.mark.parametrize("field,value,message", [
    ("total_trolleys", "-1", "Total Trolley must be a whole number from 0 to 9999."),
    ("total_trolleys", "2.5", "Total Trolley must be a whole number from 0 to 9999."),
    ("empty_trolleys", "abc", "Empty Trolley must be a whole number from 0 to 9999."),
    ("empty_trolleys", "10000", "Empty Trolley must be a whole number from 0 to 9999."),
])
def test_invalid_count_rejected(app, auth_client, field, value, message):
    r = create_part(auth_client, **{field: value})
    assert r.status_code == 400
    assert message.encode() in r.data
    assert db_rows(app, "SELECT COUNT(*) FROM parts")[0][0] == 0


def test_empty_cannot_exceed_total(app, auth_client):
    r = create_part(auth_client, total_trolleys="3", empty_trolleys="4")
    assert r.status_code == 400
    assert b"Empty Trolley cannot be more than Total Trolley." in r.data
    assert db_rows(app, "SELECT COUNT(*) FROM parts")[0][0] == 0


def test_form_fields_follow_current_storage(auth_client):
    html = auth_client.get("/parts/new").get_data(as_text=True)
    ols, total, empty = (html.index(f'name="{n}"') for n in ("ols_storage", "total_trolleys", "empty_trolleys"))
    assert ols < total < empty
    assert "Total Trolley" in html and "Empty Trolley" in html


def test_edit_changes_and_clears_counts(app, auth_client):
    create_part(auth_client, total_trolleys="6", empty_trolleys="2")
    form = auth_client.get("/parts/1/edit").get_data(as_text=True)
    assert 'name="total_trolleys" value="6"' in form and 'name="empty_trolleys" value="2"' in form
    data = {"part_name": "Bolt 1", "part_number": "PN-0001", "line_storage": "L-1", "ols_storage": "O-1"}
    auth_client.post("/parts/1/update", data={**data, "total_trolleys": "9", "empty_trolleys": "0"})
    assert counts(app) == (9, 0)
    auth_client.post("/parts/1/update", data=data)
    assert counts(app) == (None, None)


def test_dashboard_and_details_show_counts(auth_client):
    create_part(auth_client, 1, total_trolleys="10", empty_trolleys="3")
    create_part(auth_client, 2, total_trolleys="5", empty_trolleys="0")
    create_part(auth_client, 3)
    html = auth_client.get("/").get_data(as_text=True)
    assert 'id="total-trolleys">15<' in html and 'id="empty-trolleys">3<' in html
    row = html[html.index('data-label="Current Storage"'):]
    assert row.index('data-label="Total Trolley" class="cell-count">10<') < row.index(
        'data-label="Empty Trolley" class="cell-count">3<')
    details = auth_client.get("/parts/1").get_data(as_text=True)
    assert "<dt>Total Trolley</dt><dd>10</dd>" in details
    assert "<dt>Empty Trolley</dt><dd>3</dd>" in details


def test_excel_export_has_count_columns(auth_client):
    create_part(auth_client, total_trolleys="7", empty_trolleys="2")
    r = auth_client.get("/parts/export?format=xlsx")
    ws = load_workbook(io.BytesIO(r.data)).active
    assert [c.value for c in ws[4]][-3:] == ["Current Storage", "Total Trolley", "Empty Trolley"]
    assert [c.value for c in ws[5]][-2:] == [7, 2]


def test_database_rejects_negative_count(app):
    from database import IntegrityError, connect, database_target

    conn = connect(database_target(app.config))
    with pytest.raises(IntegrityError):
        conn.execute(
            "INSERT INTO parts (part_name, part_number, line_storage, ols_storage, total_trolleys, "
            "created_at, updated_at) VALUES ('a', 'b', 'c', 'd', -1, 'x', 'x')"
        )
    conn.close()


@pytest.mark.sqlite_only
def test_existing_database_gets_count_columns(tmp_path):
    cfg = make_config(tmp_path)
    conn = sqlite3.connect(cfg["DATABASE"])
    conn.execute(
        "CREATE TABLE parts (id INTEGER PRIMARY KEY AUTOINCREMENT, part_name TEXT NOT NULL, "
        "part_number TEXT NOT NULL, line_storage TEXT NOT NULL, ols_storage TEXT NOT NULL, "
        "image_filename TEXT, trolley_qty INTEGER, created_at TEXT NOT NULL, updated_at TEXT NOT NULL)"
    )
    conn.execute("INSERT INTO parts VALUES (1, 'Bolt', 'PN', 'L', 'O', NULL, 20, 'c', 'u')")
    conn.commit()
    conn.close()

    app = create_app(cfg)
    rows = db_rows(app, "SELECT part_name, trolley_qty, total_trolleys, empty_trolleys FROM parts")
    assert [tuple(r) for r in rows] == [("Bolt", 20, None, None)]
    create_app(cfg)  # running the migration again is harmless
