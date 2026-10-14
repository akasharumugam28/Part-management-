import io

import pytest
from openpyxl import load_workbook

from conftest import create_part, db_rows, upload, upload_files


def add_sub(client, part_id=1, i=1, files=None, **overrides):
    data = {"part_name": f"Washer {i}", "part_number": f"SP-{i:03d}", "trolley_qty": ""}
    data.update(overrides)
    if files:
        data.update(files)
    return client.post(f"/parts/{part_id}/sub-parts/create", data=data, content_type="multipart/form-data")


def subs(app):
    return db_rows(app, "SELECT * FROM sub_parts ORDER BY id")


@pytest.fixture
def part(auth_client):
    create_part(auth_client, 1)
    return auth_client


def test_add_sub_parts_inside_a_part(app, part):
    page = part.get("/parts/1").get_data(as_text=True)
    assert "Sub-parts" in page and "No sub-parts yet" in page and "/parts/1/sub-parts/new" in page
    form = part.get("/parts/1/sub-parts/new").get_data(as_text=True)
    assert 'name="part_name"' in form and 'data-qty-preset="40"' in form and "Bolt 1" in form

    r = add_sub(part, i=1, trolley_qty="20", files={"image": upload()})
    assert r.status_code == 303 and r.location == "/parts/1"
    r = add_sub(part, i=2, add_another="1")
    assert r.location == "/parts/1/sub-parts/new"
    rows = subs(app)
    assert [(s["part_id"], s["part_number"], s["trolley_qty"]) for s in rows] == [(1, "SP-001", 20), (1, "SP-002", None)]
    assert rows[0]["image_filename"] in upload_files(app)

    page = part.get("/parts/1").get_data(as_text=True)
    assert "Washer 1" in page and "SP-002" in page and 'id="sub-count">2<' in page
    dash = part.get("/").get_data(as_text=True)
    assert "2 sub-parts" in dash and 'id="total-sub-parts">2<' in dash


def test_sub_part_validation(app, part):
    r = add_sub(part, part_name="")
    assert r.status_code == 400 and b"Please fill in all required fields." in r.data
    r = add_sub(part, trolley_qty="0")
    assert r.status_code == 400 and b"Trolley Qty must be a whole number" in r.data
    assert subs(app) == []


def test_edit_and_delete_sub_part(app, part):
    add_sub(part, files={"image": upload()})
    old_image = subs(app)[0]["image_filename"]
    assert 'value="Washer 1"' in part.get("/parts/1/sub-parts/1/edit").get_data(as_text=True)
    r = part.post("/parts/1/sub-parts/1/update",
                  data={"part_name": "Spring", "part_number": "SP-9", "trolley_qty": "40", "remove_image": "1"})
    assert r.status_code == 303 and r.location == "/parts/1"
    row = subs(app)[0]
    assert (row["part_name"], row["trolley_qty"], row["image_filename"]) == ("Spring", 40, None)
    assert old_image not in upload_files(app)

    r = part.post("/parts/1/sub-parts/1/delete")
    assert r.status_code == 303 and subs(app) == []
    assert db_rows(app, "SELECT COUNT(*) FROM parts")[0][0] == 1  # the part stays


def test_sub_part_must_belong_to_part(part):
    create_part(part, 2)
    add_sub(part, part_id=1)
    assert part.get("/parts/2/sub-parts/1/edit").status_code == 404
    assert part.post("/parts/2/sub-parts/1/delete").status_code == 404
    assert part.get("/parts/99/sub-parts/new").status_code == 404
    assert add_sub(part, part_id=99).status_code == 404


def test_deleting_part_removes_its_sub_parts_and_images(app, part):
    add_sub(part, i=1, files={"image": upload()})
    add_sub(part, i=2, files={"image": upload()})
    assert len(upload_files(app)) == 2
    part.post("/parts/1/delete")
    assert subs(app) == [] and upload_files(app) == []


def test_search_finds_part_by_sub_part(part):
    create_part(part, 2)
    add_sub(part, part_id=2, part_number="UNIQUE-77")
    page = part.get("/?q=unique-77").get_data(as_text=True)
    assert "PN-0002" in page and "PN-0001" not in page


def test_cleanup_keeps_sub_part_images(app, part):
    add_sub(part, files={"image": upload()})
    image = subs(app)[0]["image_filename"]
    result = app.test_cli_runner().invoke(args=["cleanup-uploads"])
    assert "0 unreferenced" in result.output and image in upload_files(app)


def test_download_all(app, part):
    create_part(part, 2)
    add_sub(part, part_id=1, i=1, trolley_qty="20")
    add_sub(part, part_id=1, i=2)
    r = part.get("/parts/export?format=xlsx&name=Everything")
    assert r.status_code == 200
    ws = load_workbook(io.BytesIO(r.data)).active
    assert "2 parts, 2 sub-parts" in ws["A2"].value
    col_a = [ws.cell(row, 1).value for row in range(5, 9)]
    col_e = [ws.cell(row, 5).value for row in range(5, 9)]
    assert col_a == ["1", "1.1", "1.2", "2"]
    assert col_e == ["PN-0001", "SP-001", "SP-002", "PN-0002"]
    assert ws["C6"].value == 20 and ws["D6"].value == "↳ Washer 1"
    pdf = part.get("/parts/export?format=pdf&name=Everything")
    assert pdf.status_code == 200 and pdf.data.startswith(b"%PDF")


def test_single_part_download_includes_its_sub_parts(part):
    create_part(part, 2)
    add_sub(part, part_id=1)
    add_sub(part, part_id=2, i=5)
    ws = load_workbook(io.BytesIO(part.get("/parts/1/export?format=xlsx&name=x").data)).active
    assert [ws.cell(r, 5).value for r in range(5, 8)] == ["PN-0001", "SP-001", None]


def test_download_all_requires_login(client):
    r = client.get("/parts/export?format=pdf&name=x")
    assert r.status_code == 302 and "/login" in r.location
