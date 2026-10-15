"""Excel and PDF exports of parts and their sub-parts."""
import io
import logging
import re
from datetime import datetime
from xml.sax.saxutils import escape

from openpyxl import Workbook
from openpyxl.drawing.image import Image as XlImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from PIL import Image, UnidentifiedImageError
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Image as PdfImage
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

import images

log = logging.getLogger(__name__)

FORMATS = {
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "pdf": "application/pdf",
}
DEFAULT_NAME = "parts-list"
HEADERS = ("S.No", "Image", "Trolley Qty", "Part Name", "Part Number", "Line Storage", "Current Storage")
TEXT_FIELDS = ("part_name", "part_number", "line_storage", "ols_storage")
THUMB_PX = 60
TEAL = "0F766E"
SUB_FILL = "F0FDFA"


def clean_filename(name):
    """Keep a user-entered file name safe: letters, digits, space, - _ . only."""
    name = re.sub(r"\.(xlsx|pdf)$", "", (name or "").strip(), flags=re.I)
    name = re.sub(r"[^\w\- .]", "", name).strip(" .")[:80]
    return name or DEFAULT_NAME


def _rows(parts, subs):
    """Flatten parts and their sub-parts into rows: (sno, record, is_sub).

    Sub-parts follow their part and are numbered 1.1, 1.2, ...
    """
    for i, p in enumerate(parts, start=1):
        yield str(i), p, False
        for j, s in enumerate(subs.get(p["id"], ()), start=1):
            yield f"{i}.{j}", s, True


def _value(record, field):
    # Sub-parts have no storage columns of their own.
    return record[field] if field in record.keys() else ""


def _summary(parts, subs):
    n_parts = len(parts)
    n_subs = sum(len(subs.get(p["id"], ())) for p in parts)
    text = f"{n_parts} part{'' if n_parts == 1 else 's'}"
    if n_subs:
        text += f", {n_subs} sub-part{'' if n_subs == 1 else 's'}"
    return f"Generated {datetime.now():%d %b %Y %H:%M} · {text}"


def _thumbnails(parts, subs):
    """Map image filename -> small PNG bytes (or None), fetched in one go."""
    names = [rec["image_filename"] for _, rec, _ in _rows(parts, subs)]
    return {name: _thumbnail(name, data) for name, data in images.read_many(names).items()}


def _thumbnail(filename, data):
    """Return a small PNG of the image as bytes, or None."""
    if data is None:
        log.warning("Could not read image %r for export", filename)
        return None
    try:
        with Image.open(io.BytesIO(data)) as img:
            img = img.convert("RGB")
            img.thumbnail((THUMB_PX * 2, THUMB_PX * 2))
            buf = io.BytesIO()
            img.save(buf, "PNG")
            return buf.getvalue()
    except (OSError, UnidentifiedImageError):
        log.warning("Could not read image %r for export", filename)
        return None


def _scaled(png, box):
    with Image.open(io.BytesIO(png)) as img:
        w, h = img.size
    scale = min(box / w, box / h)
    return w * scale, h * scale


def build_xlsx(parts, subs, title):
    thumbs = _thumbnails(parts, subs)
    wb = Workbook()
    ws = wb.active
    ws.title = "Parts"
    ws.append([title])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([_summary(parts, subs)])
    ws["A2"].font = Font(italic=True, color="666666")
    ws.append([])
    ws.append(HEADERS)

    thin = Side(style="thin", color="CCCCCC")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for cell in ws[4]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=TEAL)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border

    for row, (sno, rec, is_sub) in enumerate(_rows(parts, subs), start=5):
        name = f"↳ {rec['part_name']}" if is_sub else rec["part_name"]
        ws.append([sno, "", rec["trolley_qty"], name, *(_value(rec, f) for f in TEXT_FIELDS[1:])])
        ws.row_dimensions[row].height = THUMB_PX * 0.75 + 6  # points
        for cell in ws[row]:
            indent = 2 if is_sub and cell.column == 4 else 0
            cell.alignment = Alignment(vertical="center", wrap_text=True, indent=indent)
            cell.border = border
            if is_sub:
                cell.fill = PatternFill("solid", fgColor=SUB_FILL)
            elif cell.column == 4:
                cell.font = Font(bold=True)
        for col in (1, 3):
            ws.cell(row, col).alignment = Alignment(horizontal="center", vertical="center")
        png = thumbs.get(rec["image_filename"])
        if png:
            pic = XlImage(io.BytesIO(png))
            pic.width, pic.height = _scaled(png, THUMB_PX)
            ws.add_image(pic, f"B{row}")

    for col, width in zip("ABCDEFG", (7, 11, 12, 34, 22, 18, 18)):
        ws.column_dimensions[col].width = width
    ws.freeze_panes = "A5"

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def build_pdf(parts, subs, title):
    thumbs = _thumbnails(parts, subs)
    out = io.BytesIO()
    doc = SimpleDocTemplate(
        out, pagesize=landscape(A4), title=title,
        leftMargin=12 * mm, rightMargin=12 * mm, topMargin=12 * mm, bottomMargin=12 * mm,
    )
    styles = getSampleStyleSheet()
    cell = styles["BodyText"]
    bold = ParagraphStyle("PartName", parent=cell, fontName="Helvetica-Bold")
    sub_cell = ParagraphStyle("SubName", parent=cell, leftIndent=4 * mm)
    story = [
        Paragraph(escape(title), styles["Title"]),
        Paragraph(_summary(parts, subs), styles["Normal"]),
        Spacer(1, 6 * mm),
    ]

    thumb = 16 * mm
    rows = [list(HEADERS)]
    sub_rows = []
    for sno, rec, is_sub in _rows(parts, subs):
        png = thumbs.get(rec["image_filename"])
        if png:
            w, h = _scaled(png, thumb)
            img = PdfImage(io.BytesIO(png), width=w, height=h)
        else:
            img = "-"
        name = Paragraph(("› " if is_sub else "") + escape(rec["part_name"]), sub_cell if is_sub else bold)
        rows.append([
            sno, img, str(rec["trolley_qty"] or "-"), name,
            *(Paragraph(escape(_value(rec, f)), cell) for f in TEXT_FIELDS[1:]),
        ])
        if is_sub:
            sub_rows.append(len(rows) - 1)
    if len(rows) == 1:
        rows.append(["", "", "", Paragraph("No parts", cell), "", "", ""])

    style = [
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#" + TEAL)),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTNAME", (2, 1), (2, -1), "Helvetica-Bold"),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#CCCCCC")),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("ALIGN", (0, 0), (2, -1), "CENTER"),
    ]
    style += [("BACKGROUND", (0, r), (-1, r), colors.HexColor("#" + SUB_FILL)) for r in sub_rows]
    table = Table(rows, colWidths=[14 * mm, 22 * mm, 22 * mm, 71 * mm, 52 * mm, 46 * mm, 46 * mm], repeatRows=1)
    table.setStyle(TableStyle(style))
    story.append(table)
    doc.build(story)
    return out.getvalue()


def build(fmt, parts, subs, title):
    """``subs`` maps part id -> list of that part's sub-parts."""
    return build_xlsx(parts, subs, title) if fmt == "xlsx" else build_pdf(parts, subs, title)
